#!/usr/bin/env python3
"""Generate .docx deliverables for the Future Builders Live mega-eval run."""

import os
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))


def read_md(filename):
    with open(os.path.join(BASE, filename), encoding="utf-8") as f:
        return f.read()


def style_heading(para, level=1):
    colors = {1: "1A3C6E", 2: "2E5FA3", 3: "4A7FC1"}
    run = para.runs[0] if para.runs else para.add_run(para.text)
    run.bold = True
    if level in colors:
        run.font.color.rgb = RGBColor.from_string(colors[level])
    run.font.size = Pt({1: 18, 2: 14, 3: 12}.get(level, 11))


def add_title_block(doc, title, subtitle=None):
    doc.add_paragraph()
    t = doc.add_heading(title, level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in t.runs:
        run.font.color.rgb = RGBColor.from_string("1A3C6E")
        run.font.size = Pt(22)
    if subtitle:
        s = doc.add_paragraph(subtitle)
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in s.runs:
            run.font.color.rgb = RGBColor.from_string("555555")
    doc.add_paragraph()


def add_md_content(doc, md_text, skip_first_h1=True):
    """Parse markdown and add to doc with basic formatting."""
    lines = md_text.split("\n")
    i = 0
    first_h1_seen = False
    in_table = False
    table_rows = []

    while i < len(lines):
        line = lines[i]

        # Detect table
        if line.startswith("|"):
            table_rows.append(line)
            i += 1
            continue
        else:
            if table_rows:
                flush_table(doc, table_rows)
                table_rows = []

        # H1
        if line.startswith("# ") and not line.startswith("## "):
            if skip_first_h1 and not first_h1_seen:
                first_h1_seen = True
                i += 1
                continue
            first_h1_seen = True
            h = doc.add_heading(line[2:].strip(), level=1)
            style_heading(h, 1)

        # H2
        elif line.startswith("## "):
            h = doc.add_heading(line[3:].strip(), level=2)
            style_heading(h, 2)

        # H3
        elif line.startswith("### "):
            h = doc.add_heading(line[4:].strip(), level=3)
            style_heading(h, 3)

        # Horizontal rule
        elif line.strip() in ("---", "***", "___"):
            doc.add_paragraph()

        # Bullet
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_formatting(p, line[2:].strip())

        # Numbered list
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            content = re.sub(r"^\d+\. ", "", line)
            add_inline_formatting(p, content)

        # Empty line
        elif line.strip() == "":
            pass

        # Normal paragraph
        else:
            p = doc.add_paragraph()
            add_inline_formatting(p, line.strip())

        i += 1

    if table_rows:
        flush_table(doc, table_rows)


def flush_table(doc, rows):
    """Render markdown table rows into a docx table."""
    data_rows = [r for r in rows if not re.match(r"^\|[-| :]+\|$", r.strip())]
    if not data_rows:
        return
    parsed = []
    for row in data_rows:
        cells = [c.strip() for c in row.strip().strip("|").split("|")]
        parsed.append(cells)
    if not parsed:
        return
    max_cols = max(len(r) for r in parsed)
    t = doc.add_table(rows=len(parsed), cols=max_cols)
    t.style = "Table Grid"
    for ri, row in enumerate(parsed):
        for ci, cell_text in enumerate(row):
            cell = t.cell(ri, ci)
            p = cell.paragraphs[0]
            add_inline_formatting(p, cell_text)
            if ri == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()


def add_inline_formatting(para, text):
    """Handle **bold**, _italic_, and plain text inline."""
    pattern = re.compile(r"(\*\*[^*]+\*\*|_[^_]+_|\*[^*]+\*|`[^`]+`)")
    parts = pattern.split(text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = para.add_run(part[2:-2])
            run.bold = True
        elif (part.startswith("_") and part.endswith("_")) or (
            part.startswith("*") and part.endswith("*") and not part.startswith("**")
        ):
            run = para.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = para.add_run(part[1:-1])
            run.font.name = "Courier New"
        else:
            para.add_run(part)


# ─── Executive Summary ────────────────────────────────────────────────────────

def build_executive_summary():
    doc = Document()
    add_title_block(
        doc,
        "Future Builders Live",
        "Mega-Eval Executive Summary  •  May 2026",
    )

    content = """
## Verdict

Future Builders Live is a compelling concept with genuine market timing: AI/robotics interest is at a cultural peak, premium family experience demand is growing, and no single competitor has packaged the "TED-style talk + maker demo + student showcase" format for families in a consistent, high-quality series. However, the pitch is incomplete as a business plan—it is missing ticket prices, founder credibility, a full cost model, and a student presenter recruitment system. The execution risk of consistently delivering premium live events is substantial. This is a fundable and buildable concept, but it requires significant operational infrastructure before the first event or it will underdeliver on its premium promise.

## Top 3 Strengths

- **Powerful emotional core:** Kids presenting their real work to a real audience is genuinely differentiated from every identified competitor
- **AI/robotics market timing:** The cultural moment for this content is historically strong in 2025–2026
- **Format clarity:** The defined event structure makes it repeatable and programmable vs. open-ended alternatives

## Top 3 Risks

- **Premium-execution gap:** One underfunded, chaotic, or poorly attended event permanently damages the premium brand in that community
- **Sponsor revenue volatility:** Local sponsorship is the hardest, most labor-intensive, most churn-prone revenue stream
- **Founder anonymity:** No social proof of who is running this or why they can be trusted with a child-facing, paid public event

## Top 3 Opportunities

- **Content capture:** Recording student talks and expert presentations creates viral-potential media at zero marginal cost
- **Workshop-to-membership flywheel:** Post-event conversion to recurring revenue is the highest-value lever for financial sustainability
- **City licensing model:** A documented event playbook enables city-to-city expansion without requiring the founder at every event

## Priority Actions (urgent first)

1. State a ticket price and publish a landing page with email capture — validate demand before any production investment
2. Write and publish a founder story with real name, photo, and relevant background
3. Build a 2-page sponsor prospectus with 3 tiers and send to 10 local businesses this week
4. Write a Student Presenter Brief and identify one school or community group as a recruitment beta partner
5. Build a full cost + labor pro forma (including founder time at market rate) for Event 1
"""
    add_md_content(doc, content, skip_first_h1=False)
    path = os.path.join(BASE, "00-executive-summary.docx")
    doc.save(path)
    print(f"  ✓ {path}")


# ─── Hater Mode ───────────────────────────────────────────────────────────────

def build_hater_mode():
    doc = Document()
    add_title_block(
        doc,
        "Future Builders Live — Hater Mode Feedback",
        "Phase 1A: 12-Persona Critical Analysis  •  May 2026",
    )
    md = read_md("phase1a-hater-raw.md")
    add_md_content(doc, md)
    path = os.path.join(BASE, "01-hater-mode-feedback.docx")
    doc.save(path)
    print(f"  ✓ {path}")


# ─── Competitive Landscape ────────────────────────────────────────────────────

def build_competitive():
    doc = Document()
    add_title_block(
        doc,
        "Future Builders Live — Competitive Landscape",
        "Phase 1B: Market Analysis  •  May 2026",
    )
    md = read_md("phase1b-competitive-raw.md")
    add_md_content(doc, md)
    path = os.path.join(BASE, "02-competitive-landscape.docx")
    doc.save(path)
    print(f"  ✓ {path}")


# ─── Strengths & Opportunities ────────────────────────────────────────────────

def build_strengths():
    doc = Document()
    add_title_block(
        doc,
        "Future Builders Live — Strengths & Opportunities",
        "Phase 1C: What's Working & Where to Grow  •  May 2026",
    )
    md = read_md("phase1c-strengths-raw.md")
    add_md_content(doc, md)
    path = os.path.join(BASE, "03-strengths-opportunities.docx")
    doc.save(path)
    print(f"  ✓ {path}")


# ─── Critical Fixes & Design ──────────────────────────────────────────────────

def build_synthesis():
    doc = Document()
    add_title_block(
        doc,
        "Future Builders Live — Critical Fixes & Design",
        "Phase 2: Synthesis  •  May 2026",
    )
    md = read_md("phase2-synthesis.md")
    add_md_content(doc, md)
    path = os.path.join(BASE, "04-critical-fixes-and-design.docx")
    doc.save(path)
    print(f"  ✓ {path}")


# ─── Content Strategy Outline ─────────────────────────────────────────────────

def build_content_strategy():
    doc = Document()
    add_title_block(
        doc,
        "Future Builders Live — Content Strategy Outline",
        "Phase 3: Publishable Content Plan  •  May 2026",
    )
    md = read_md("phase3-content-outline-raw.md")
    add_md_content(doc, md)
    path = os.path.join(BASE, "05-content-strategy-outline.docx")
    doc.save(path)
    print(f"  ✓ {path}")


if __name__ == "__main__":
    print("Generating Future Builders Live mega-eval deliverables...")
    build_executive_summary()
    build_hater_mode()
    build_competitive()
    build_strengths()
    build_synthesis()
    build_content_strategy()
    print("Done. 6 .docx files written.")
